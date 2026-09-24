"""Собирает tco_model.xlsx с живыми формулами из параметров tco_model.py.

Запуск: python finance/generate_xlsx.py
Файл открывается в LibreOffice Calc, Google Sheets и Excel; любой параметр на листе
«Параметры» можно поменять — пересчитаются все листы.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from tco_model import ASSUMPTIONS, SCENARIOS, SOURCES, P

OUT = Path(__file__).with_name("tco_model.xlsx")
YEARS = ["Год 0 (внедрение)", "Год 1", "Год 2", "Год 3"]
COLS = "CDEF"          # годы 0..3
TOTAL = "G"

HEAD = PatternFill("solid", fgColor="1F4E79")
SUB = PatternFill("solid", fgColor="DDEBF7")
INPUT = PatternFill("solid", fgColor="FFF2CC")
BOLD = Font(bold=True)
WHITE = Font(bold=True, color="FFFFFF")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
MONEY = '#,##0" ₸"'
PCT = "0%"

PARAM_META = {
    "usd_kzt": ("Курс USD/KZT", "₸", "S6"),
    "discount_rate": ("Ставка дисконтирования (базовая ставка НБ РК)", "%", "S7"),
    "vat": ("НДС", "%", "S8"),
    "salary_edu": ("Средняя зарплата в образовании (методист, куратор)", "₸/мес", "S1"),
    "salary_dev": ("Средняя зарплата Junior Python-разработчика", "₸/мес", "S2"),
    "payroll_factor": ("Коэффициент начислений на ФОТ", "×", "A1"),
    "hours_month": ("Рабочих часов в месяце", "ч", "A2"),
    "methodist_h_before": ("Методист: сведение заявок сейчас", "ч/нед", "A3"),
    "methodist_h_after": ("Методист: сведение заявок после внедрения", "ч/нед", "A3"),
    "active_weeks": ("Недель в году с потоком заявок", "нед", "A4"),
    "curators": ("Количество кураторов практики", "чел", "A5"),
    "curator_h_week": ("Куратор: выяснение статусов и сроков", "ч/нед", "A5"),
    "curator_saving": ("Доля времени кураторов, снимаемая системой", "%", "A5"),
    "curator_weeks": ("Недель практики с отчётами", "нед", "A5"),
    "lost_cases_year": ("Потерянных заявок / жалоб в год", "шт", "A16"),
    "lost_case_hours": ("Часов на разбор одной потерянной заявки", "ч", "A16"),
    "ramp_y1": ("Реализация эффекта в год 1", "%", "A6"),
    "dev_hours": ("Трудоёмкость разработки MVP", "ч", "A7"),
    "pm_ba_hours": ("Трудоёмкость PM и аналитики", "ч", "A7"),
    "migration_hours": ("Перенос данных из Excel", "ч", "A8"),
    "users_trained": ("Пользователей к обучению", "чел", "A9"),
    "training_h_user": ("Часов обучения на пользователя", "ч", "A9"),
    "retraining_h_year": ("Обучение новых кураторов", "ч/год", "A9"),
    "support_h_month": ("Поддержка своей системы", "ч/мес", "A10"),
    "vps_month": ("VPS hoster.kz Cloud 1-2-50 (прод)", "₸/мес", "S3"),
    "backup_vps_month": ("VPS hoster.kz Cloud 1-1-25 (резервные копии)", "₸/мес", "S3"),
    "ws_license_usd": ("Windows Server 2025 Standard, 16 ядер", "USD", "S5"),
    "sql_2core_usd": ("SQL Server Standard, пакет 2 ядра", "USD", "S5"),
    "sql_packs": ("Пакетов SQL Server (минимум 4 ядра)", "шт", "S5"),
    "b24_month": ("Битрикс24 «Стандартный» (оплата за год, без НДС)", "₸/мес", "S4"),
    "b24_setup_hours": ("Битрикс24: настройка", "ч", "A11"),
    "b24_support_h_month": ("Битрикс24: администрирование", "ч/мес", "A11"),
    "b24_benefit_share": ("Битрикс24: доля закрываемого эффекта", "%", "A11"),
    "decommission_hours": ("Вывод своей системы из эксплуатации", "ч", "A12"),
    "b24_decommission_hours": ("Выгрузка данных из Битрикс24", "ч", "A12"),
}


def name(wb: Workbook, key: str, ref: str) -> None:
    wb.defined_names[key] = DefinedName(key, attr_text=ref)


def header(ws, row: int, values: list[str]) -> None:
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=i + 1, value=v)
        c.fill, c.font = HEAD, WHITE
        c.alignment = Alignment(wrap_text=True, vertical="center")


def sheet_params(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "Параметры"
    ws["A1"] = "Параметры модели. Жёлтые ячейки можно менять — пересчитаются все листы."
    ws["A1"].font = BOLD
    header(ws, 3, ["Ключ", "Параметр", "Значение", "Ед.", "Источник / допущение"])
    for i, key in enumerate(PARAM_META, start=4):
        label, unit, src = PARAM_META[key]
        ws.cell(row=i, column=1, value=key)
        ws.cell(row=i, column=2, value=label)
        c = ws.cell(row=i, column=3, value=P[key])
        c.fill = INPUT
        c.number_format = PCT if unit == "%" else "#,##0.##"
        ws.cell(row=i, column=4, value=unit)
        ws.cell(row=i, column=5, value=src)
        name(wb, key, f"'Параметры'!$C${i}")
    for col, w in zip("ABCDE", (22, 52, 14, 8, 22)):
        ws.column_dimensions[col].width = w


def sheet_benefit(wb: Workbook) -> None:
    ws = wb.create_sheet("Эффект")
    ws["A1"] = "Ставки и годовой эффект (формулы видны в ячейках столбца B)"
    ws["A1"].font = BOLD
    rows = [
        ("rate_edu", "Стоимость часа методиста/куратора, ₸", "=salary_edu*payroll_factor/hours_month", MONEY),
        ("rate_dev", "Стоимость часа разработчика, ₸", "=salary_dev*payroll_factor/hours_month", MONEY),
        ("h_methodist", "Сэкономлено часов методиста в год",
         "=(methodist_h_before-methodist_h_after)*active_weeks", "0"),
        ("h_curators", "Сэкономлено часов кураторов в год",
         "=curators*curator_h_week*curator_saving*curator_weeks", "0"),
        ("h_lost", "Часы на разбор потерянных заявок в год", "=lost_cases_year*lost_case_hours", "0"),
        ("h_total", "Итого сэкономлено часов в год", "=h_methodist+h_curators+h_lost", "0"),
        ("annual_benefit", "Годовой эффект при 100 % реализации, ₸", "=h_total*rate_edu", MONEY),
        ("cost_of_delay", "Cost of delay: потеря в месяц без системы, ₸", "=annual_benefit/12", MONEY),
    ]
    for i, (key, label, formula, fmt) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=label)
        c = ws.cell(row=i, column=2, value=formula)
        c.number_format = fmt
        name(wb, key, f"'Эффект'!$B${i}")
    ws["A12"] = ("Не монетизировано: прослеживаемый статус заявки (0 → 100 %), снижение риска по закону РК "
                 "о персональных данных (заявки с ПДн сейчас лежат в почте и Excel), время студентов.")
    ws["A12"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A12:B14")
    ws.column_dimensions["A"].width = 55
    ws.column_dimensions["B"].width = 20


# Строка затрат: подпись + формулы по годам 0..3 (None = 0)
def own_build_lines(dev_paid: bool, licenses: bool) -> list[tuple[str, list[str | None]]]:
    infra = "(vps_month+backup_vps_month)*12"
    run_train = "retraining_h_year*rate_edu"
    support = "support_h_month*12*rate_dev"
    lines = [
        ("Разработка MVP" + ("" if dev_paid else " (труд студентов, A13)"),
         ["(dev_hours+pm_ba_hours)*rate_dev" if dev_paid else "0", None, None, None]),
    ]
    if licenses:
        lines.append(("Лицензии Windows Server + SQL Server (с НДС)",
                      ["(ws_license_usd+sql_2core_usd*sql_packs)*usd_kzt*(1+vat)", None, None, None]))
    lines += [
        ("Инфраструктура: VPS + резервные копии", [f"{infra}/3", infra, infra, infra]),
        ("Перенос данных из Excel", ["migration_hours*rate_dev", None, None, None]),
        ("Обучение персонала", ["users_trained*training_h_user*rate_edu", run_train, run_train, run_train]),
        ("Поддержка и сопровождение", [None, support, support, support]),
        ("Вывод из эксплуатации", [None, None, None, "decommission_hours*rate_dev"]),
    ]
    return lines


def b24_lines() -> list[tuple[str, list[str | None]]]:
    sub = "b24_month*12*(1+vat)"
    run_train = "retraining_h_year*rate_edu"
    adm = "b24_support_h_month*12*rate_dev"
    return [
        ("Подписка Битрикс24 «Стандартный» (с НДС)", [f"{sub}/3", sub, sub, sub]),
        ("Настройка портала", ["b24_setup_hours*rate_dev", None, None, None]),
        ("Перенос данных из Excel", ["migration_hours*rate_dev", None, None, None]),
        ("Обучение персонала", ["users_trained*training_h_user*rate_edu", run_train, run_train, run_train]),
        ("Администрирование", [None, adm, adm, adm]),
        ("Выгрузка данных при отказе", [None, None, None, "b24_decommission_hours*rate_dev"]),
    ]


OPTIONS = [
    ("V0", "В0. Ничего не делать (Excel + почта)", [], "0"),
    ("V1", "В1. Купить готовое: Битрикс24 «Стандартный»", b24_lines(), "b24_benefit_share"),
    ("V2", "В2. Своя разработка, проприетарный стек (Windows Server + SQL Server)",
     own_build_lines(True, True), "1"),
    ("V3", "В3. Своя разработка, открытый стек (Linux + PostgreSQL + Python)",
     own_build_lines(True, False), "1"),
    ("V3S", "В3-С. Открытый стек силами студенческой команды (рекомендуется)",
     own_build_lines(False, False), "1"),
]


def sheet_tco(wb: Workbook) -> dict[str, dict[str, int]]:
    ws = wb.create_sheet("TCO")
    ws["A1"] = "TCO на 3 года по вариантам, ₸. Год 0 — внедрение (4 месяца), годы 1–3 — эксплуатация."
    ws["A1"].font = BOLD
    header(ws, 3, ["Вариант / статья", "", *YEARS, "Итого"])
    ws["B4"] = "Коэф. дисконтирования"
    for j, col in enumerate(COLS):
        ws[f"{col}4"] = f"=1/(1+discount_rate)^{j}"
        ws[f"{col}4"].number_format = "0.000"
    name(wb, "df_row", "'TCO'!$C$4:$F$4")

    rows: dict[str, dict[str, int]] = {}
    r = 6
    for key, title, lines, share in OPTIONS:
        ws.cell(row=r, column=1, value=title).font = BOLD
        for c in range(1, 8):
            ws.cell(row=r, column=c).fill = SUB
        r += 1
        first = r
        for label, formulas in lines:
            ws.cell(row=r, column=1, value="   " + label)
            for col, f in zip(COLS, formulas):
                ws[f"{col}{r}"] = f"={f}" if f else 0
                ws[f"{col}{r}"].number_format = MONEY
            ws[f"{TOTAL}{r}"] = f"=SUM(C{r}:F{r})"
            ws[f"{TOTAL}{r}"].number_format = MONEY
            r += 1
        cost = r
        ws.cell(row=r, column=1, value="Итого затраты").font = BOLD
        for col in COLS + TOTAL:
            ws[f"{col}{r}"] = f"=SUM({col}{first}:{col}{r - 1})" if lines else 0
            ws[f"{col}{r}"].number_format = MONEY
            ws[f"{col}{r}"].font = BOLD
        r += 1
        ben = r
        ws.cell(row=r, column=1, value="Выгода (экономия времени персонала)")
        for j, col in enumerate(COLS):
            k = ["0", "ramp_y1", "1", "1"][j]
            ws[f"{col}{r}"] = f"=annual_benefit*{share}*{k}"
            ws[f"{col}{r}"].number_format = MONEY
        ws[f"{TOTAL}{r}"] = f"=SUM(C{r}:F{r})"
        ws[f"{TOTAL}{r}"].number_format = MONEY
        r += 1
        cf = r
        ws.cell(row=r, column=1, value="Чистый денежный поток")
        for col in COLS + TOTAL:
            ws[f"{col}{r}"] = f"={col}{ben}-{col}{cost}"
            ws[f"{col}{r}"].number_format = MONEY
        r += 1
        cum = r
        ws.cell(row=r, column=1, value="Накопленный поток")
        ws[f"C{r}"] = f"=C{cf}"
        for prev, col in zip(COLS, COLS[1:]):
            ws[f"{col}{r}"] = f"={prev}{r}+{col}{cf}"
        for col in COLS:
            ws[f"{col}{r}"].number_format = MONEY
            ws[f"{col}{r}"].border = THIN
        rows[key] = {"cost": cost, "ben": ben, "cf": cf, "cum": cum}
        r += 2
    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 4
    for col in COLS + TOTAL:
        ws.column_dimensions[col].width = 17
    ws.freeze_panes = "C5"
    return rows


def sheet_summary(wb: Workbook, rows: dict[str, dict[str, int]]) -> None:
    ws = wb.create_sheet("Сводка", 0)
    ws["A1"] = "Сводка: TCO, ROI, NPV и окупаемость за 3 года (базовый сценарий)"
    ws["A1"].font = Font(bold=True, size=13)
    header(ws, 3, ["Вариант", "TCO 3 года, ₸", "Выгода 3 года, ₸", "ROI", "NPV, ₸", "Окупаемость, лет эксплуатации"])
    ws["A10"] = "Формулы: ROI = (Σ выгод − Σ затрат) / Σ затрат;  NPV = Σ (выгода_t − затраты_t) / (1 + r)^t, r = discount_rate;"
    ws["A11"] = "окупаемость — момент, когда накопленный поток становится ≥ 0 (линейная интерполяция внутри года)."
    for i, (key, title, _, _) in enumerate(OPTIONS, start=4):
        t = rows[key]
        c, b, cf, cum = t["cost"], t["ben"], t["cf"], t["cum"]
        ws.cell(row=i, column=1, value=title)
        ws.cell(row=i, column=2, value=f"=TCO!G{c}").number_format = MONEY
        ws.cell(row=i, column=3, value=f"=TCO!G{b}").number_format = MONEY
        ws.cell(row=i, column=4, value=f'=IF(B{i}=0,"—",(C{i}-B{i})/B{i})').number_format = PCT
        ws.cell(row=i, column=5, value=f"=SUMPRODUCT(TCO!C{cf}:F{cf},df_row)").number_format = MONEY
        ws.cell(row=i, column=6, value=(
            f'=IF(B{i}=0,"—",IF(TCO!D{cum}>=0,-TCO!C{cum}/TCO!D{cf},'
            f'IF(TCO!E{cum}>=0,1-TCO!D{cum}/TCO!E{cf},'
            f'IF(TCO!F{cum}>=0,2-TCO!E{cum}/TCO!F{cf},"не окупается за 3 года"))))'
        )).number_format = "0.0"
    ws.column_dimensions["A"].width = 66
    for col in "BCDEF":
        ws.column_dimensions[col].width = 18

    ws["A13"] = "Открытое ПО против проприетарного (В3 против В2)"
    ws["A13"].font = BOLD
    ws["A14"] = "Разница TCO за 3 года, ₸"
    ws["B14"] = "=B6-B7"
    ws["A15"] = "из них лицензии Windows Server + SQL Server (с НДС), ₸"
    ws["B15"] = "=(ws_license_usd+sql_2core_usd*sql_packs)*usd_kzt*(1+vat)"
    ws["A16"] = "Во сколько раз лицензии больше годового эффекта"
    ws["B16"] = "=B15/annual_benefit"
    for cell in ("B14", "B15"):
        ws[cell].number_format = MONEY
    ws["B16"].number_format = "0.0"

    chart = BarChart()
    chart.type = "bar"
    chart.title = "NPV по вариантам, ₸"
    chart.add_data(Reference(ws, min_col=5, min_row=3, max_row=8), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=4, max_row=8))
    chart.legend = None
    chart.height, chart.width = 8, 22
    ws.add_chart(chart, "A19")


def sheet_sensitivity(wb: Workbook, rows: dict[str, dict[str, int]]) -> None:
    ws = wb.create_sheet("Чувствительность", 1)
    ws["A1"] = "Анализ чувствительности: затраты × k_c, выгода × k_b (жёлтые ячейки можно менять)"
    ws["A1"].font = BOLD
    ws["A3"], ws["A4"] = "k_c — множитель затрат", "k_b — множитель выгоды"
    for j, (scen, ck, bk) in enumerate(SCENARIOS):
        col = 2 + 2 * j
        ws.cell(row=2, column=col, value=scen).font = BOLD
        ws.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col + 1)
        for rr, v in ((3, ck), (4, bk)):
            c = ws.cell(row=rr, column=col, value=v)
            c.fill, c.number_format = INPUT, "0.00"
        ws.cell(row=5, column=col, value="ROI").font = BOLD
        ws.cell(row=5, column=col + 1, value="NPV, ₸").font = BOLD
    for i, (key, title, _, _) in enumerate(OPTIONS, start=6):
        t = rows[key]
        ws.cell(row=i, column=1, value=title)
        for j in range(len(SCENARIOS)):
            col = 2 + 2 * j
            L = ws.cell(row=3, column=col).column_letter
            ck, bk = f"${L}$3", f"${L}$4"
            cost, ben = f"TCO!G{t['cost']}", f"TCO!G{t['ben']}"
            ws.cell(row=i, column=col, value=f'=IF({cost}=0,"—",({bk}*{ben}-{ck}*{cost})/({ck}*{cost}))'
                    ).number_format = PCT
            ws.cell(row=i, column=col + 1, value=(
                f"={bk}*SUMPRODUCT(TCO!C{t['ben']}:F{t['ben']},df_row)"
                f"-{ck}*SUMPRODUCT(TCO!C{t['cost']}:F{t['cost']},df_row)")).number_format = MONEY
    ws["A12"] = "Точка безубыточности для В3-С: при каком множителе выгоды NPV = 0 (при базовых затратах)"
    t = rows["V3S"]
    ws["B12"] = (f"=SUMPRODUCT(TCO!C{t['cost']}:F{t['cost']},df_row)"
                 f"/SUMPRODUCT(TCO!C{t['ben']}:F{t['ben']},df_row)")
    ws["B12"].number_format = PCT
    ws.column_dimensions["A"].width = 66
    for j in range(2, 10):
        ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = 14


def sheet_sources(wb: Workbook) -> None:
    ws = wb.create_sheet("Источники и допущения")
    header(ws, 1, ["ID", "Источник (открытый, дата обращения 24.09.2026) / допущение"])
    r = 2
    for k, v in {**SOURCES, **ASSUMPTIONS}.items():
        ws.cell(row=r, column=1, value=k)
        ws.cell(row=r, column=2, value=v).alignment = Alignment(wrap_text=True)
        r += 1
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 130


def main() -> None:
    wb = Workbook()
    sheet_params(wb)
    sheet_benefit(wb)
    rows = sheet_tco(wb)
    sheet_summary(wb, rows)
    sheet_sensitivity(wb, rows)
    sheet_sources(wb)
    wb.save(OUT)
    print(f"saved {OUT}")


if __name__ == "__main__":
    main()
