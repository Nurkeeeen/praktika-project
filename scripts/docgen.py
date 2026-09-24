"""Общий генератор документов лабораторных: один список блоков → .docx и .md.

Оформление .docx повторяет ЛЗ 1–2: Times New Roman 11, заголовки 14/12 жирным, поля 2,54 см.
"""

from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


class Doc:
    def __init__(self) -> None:
        self.blocks: list[tuple] = []

    # --- содержание ---
    def title(self, t: str) -> None: self.blocks.append(("title", t))
    def sub(self, t: str) -> None: self.blocks.append(("sub", t))
    def h1(self, t: str) -> None: self.blocks.append(("h1", t))
    def h2(self, t: str) -> None: self.blocks.append(("h2", t))
    def p(self, t: str) -> None: self.blocks.append(("p", t))
    def note(self, t: str) -> None: self.blocks.append(("note", t))
    def bullets(self, items: list[str]) -> None: self.blocks.append(("ul", items))
    def numbered(self, items: list[str]) -> None: self.blocks.append(("ol", items))
    def table(self, head: list[str], rows: list[list]) -> None: self.blocks.append(("table", head, rows))

    def image(self, path: Path, caption: str, width_cm: float = 16.5) -> None:
        self.blocks.append(("image", Path(path), caption, width_cm))

    # --- рендеры ---
    def render_docx(self, path: Path) -> None:
        doc = Document()
        st = doc.styles["Normal"]
        st.font.name, st.font.size = "Times New Roman", Pt(11)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        for s in ("Heading 1", "Heading 2", "Title"):
            hs = doc.styles[s]
            hs.font.name, hs.font.bold = "Times New Roman", True
            hs.font.color.rgb = None
            hs.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        doc.styles["Heading 1"].font.size = Pt(14)
        doc.styles["Heading 2"].font.size = Pt(12)
        for sec in doc.sections:
            sec.left_margin = sec.right_margin = Cm(2.54)
        for blk in self.blocks:
            kind = blk[0]
            if kind == "title":
                r = doc.add_paragraph().add_run(blk[1])
                r.bold, r.font.size = True, Pt(14)
            elif kind in ("sub", "p"):
                doc.add_paragraph(blk[1])
            elif kind == "h1":
                doc.add_heading(blk[1], level=1)
            elif kind == "h2":
                doc.add_heading(blk[1], level=2)
            elif kind == "note":
                doc.add_paragraph().add_run(blk[1]).italic = True
            elif kind in ("ul", "ol"):
                for it in blk[1]:
                    doc.add_paragraph(it, style="List Bullet" if kind == "ul" else "List Number")
            elif kind == "table":
                _docx_table(doc, blk[1], blk[2])
            elif kind == "image":
                doc.add_picture(str(blk[1]), width=Cm(blk[3]))
                doc.add_paragraph().add_run(blk[2]).italic = True
        path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(path)

    def render_md(self, path: Path) -> None:
        out: list[str] = []
        for blk in self.blocks:
            kind = blk[0]
            if kind == "title":
                out += [f"# {blk[1]}", ""]
            elif kind == "sub":
                out += [f"_{blk[1]}_", ""]
            elif kind == "h1":
                out += [f"## {blk[1]}", ""]
            elif kind == "h2":
                out += [f"### {blk[1]}", ""]
            elif kind == "p":
                out += [blk[1], ""]
            elif kind == "note":
                out += [f"> {blk[1]}", ""]
            elif kind == "ul":
                out += [f"- {it}" for it in blk[1]] + [""]
            elif kind == "ol":
                out += [f"{i}. {it}" for i, it in enumerate(blk[1], 1)] + [""]
            elif kind == "table":
                head, rows = blk[1], blk[2]
                out += ["| " + " | ".join(_md_cell(h) for h in head) + " |", "|" + "---|" * len(head)]
                out += ["| " + " | ".join(_md_cell(v) for v in r) + " |" for r in rows]
                out.append("")
            elif kind == "image":
                rel = Path(os.path.relpath(blk[1], path.parent)).as_posix()
                out += [f"![{blk[2]}]({rel})", "", f"_{blk[2]}_", ""]
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text("\n".join(out), encoding="utf-8")


def _md_cell(v) -> str:
    return str(v).replace("|", "\\|").replace("\n", "<br>")


def _shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def _docx_table(doc, head: list[str], rows: list[list]) -> None:
    t = doc.add_table(rows=1, cols=len(head))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, h in zip(t.rows[0].cells, head, strict=True):
        c.text = ""
        c.paragraphs[0].add_run(h).bold = True
        _shade(c, "DDEBF7")
    for row in rows:
        for c, v in zip(t.add_row().cells, row, strict=True):
            c.text = str(v)
    for row in t.rows:
        for c in row.cells:
            for par in c.paragraphs:
                for run in par.runs:
                    run.font.size = Pt(9.5)
    doc.add_paragraph()
